"""Local Cross–Pollinate demo: python app.py."""
import asyncio
import json
import os
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID, uuid4

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, PlainTextResponse, Response
from fastapi.staticfiles import StaticFiles

from agent import research, codex_binary
from export_report import to_markdown
from models import ResearchRequest

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
app = FastAPI(title="Cross–Pollinate")
jobs = {}
tasks = {}


@app.get("/api/config")
def config():
    backend = os.getenv("AGENT_BACKEND", "codex")
    return {"backend": backend, "ready": bool(codex_binary()) if backend == "codex" else bool(os.getenv("OPENAI_API_KEY")),
            "name": "Cross–Pollinate"}


def run_path(run_id):
    try:
        UUID(run_id)
    except ValueError:
        raise HTTPException(404, "Run not found")
    return ROOT / "runs" / run_id


def get_report(run_id):
    path = run_path(run_id) / "report.json"
    if not path.exists():
        raise HTTPException(404, "Completed report not found")
    return json.loads(path.read_text())


async def execute(run_id, request, previous):
    job = jobs[run_id]
    def emit(kind, message, detail):
        job["events"].append({"kind": kind, "message": message, "detail": detail,
                              "time": datetime.now(timezone.utc).isoformat()})
    try:
        job["report"] = await research(request, run_path(run_id), emit, previous)
        (run_path(run_id) / "report.md").write_text(to_markdown(job["report"]))
        job["status"] = "complete"
        emit("stage", "Report ready", {})
    except asyncio.CancelledError:
        job["status"] = "cancelled"
    except Exception as error:
        # The UI must surface a failed network/model run, not silently show a canned report.
        job["status"] = "failed"
        job["error"] = str(error) or type(error).__name__
    finally:
        tasks.pop(run_id, None)


@app.post("/api/runs")
async def start(request: ResearchRequest):
    if len(tasks) >= 2:
        raise HTTPException(429, "Two runs are already active. Finish or cancel one first.")
    previous = get_report(request.previous_run) if request.previous_run else None
    run_id = str(uuid4())
    jobs[run_id] = {"id": run_id, "status": "running", "events": [], "request": request.model_dump()}
    tasks[run_id] = asyncio.create_task(execute(run_id, request, previous))
    return {"id": run_id}


@app.get("/api/runs/{run_id}")
def status(run_id: str):
    if run_id in jobs:
        return jobs[run_id]
    return {"id": run_id, "status": "complete", "events": [], "report": get_report(run_id),
            "request": json.loads((run_path(run_id) / "request.json").read_text())}


@app.post("/api/runs/{run_id}/cancel")
def cancel(run_id: str):
    if run_id in tasks:
        tasks[run_id].cancel()
    return {"cancel_requested": True}


@app.get("/api/runs/{run_id}/export/{format}")
def export(run_id: str, format: str):
    report = get_report(run_id)
    if format == "json":
        # Export citations and retrieval levels; raw abstracts stay in the local evidence log.
        for source in report.get("sources", []):
            source.pop("abstract", None)
        content, media = json.dumps(report, ensure_ascii=False, indent=2), "application/json"
    elif format == "md":
        content, media = to_markdown(report), "text/markdown"
    elif format == "docx":
        from io import BytesIO
        from docx import Document
        document = Document()
        for line in to_markdown(report).splitlines():
            if line.startswith("#"):
                depth = len(line) - len(line.lstrip("#"))
                document.add_heading(line[depth:].strip(), min(depth - 1, 3))
            elif line.startswith("- "):
                document.add_paragraph(line[2:].replace("**", ""), style="List Bullet")
            elif line:
                document.add_paragraph(line.replace("**", ""))
        buffer = BytesIO()
        document.save(buffer)
        content, media = buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise HTTPException(404, "Export format not found")
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="cross-pollinate.{format}"'})


app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


@app.get("/")
def index():
    return FileResponse(ROOT / "static" / "index.html")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=int(os.getenv("PORT", "7860")))
