"""Safe document-to-text extraction, with page/paragraph labels for traceability."""
import io
from pathlib import Path

MAX_FILE_BYTES = 3 * 1024 * 1024
ALLOWED = {'.pdf', '.docx', '.txt', '.md'}

def extract_document(filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError('Unsupported document type: use PDF, DOCX, TXT, or MD')
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('File exceeds 3 MB limit')
    if not content:
        raise ValueError('Empty file')
    if suffix == '.pdf':
        try:
            import fitz
        except ImportError as e:
            raise ValueError('PDF uploads need PyMuPDF: pip install -r requirements.txt') from e
        try:
            document = fitz.open(stream=content, filetype='pdf')
            try:
                if document.needs_pass:
                    raise ValueError('Password-protected PDF not supported')
                if len(document) > 60:
                    raise ValueError('PDF exceeds 60-page limit')
                parts=[]
                for page_no, page in enumerate(document, 1):
                    for line_no, line in enumerate(page.get_text('text').splitlines(), 1):
                        if line.strip():
                            parts.append(f'[Page {page_no}, line {line_no}] {line.strip()}')
            finally:
                document.close()
        except ValueError:
            raise
        except Exception as e:
            raise ValueError('Unable to read PDF') from e
    elif suffix == '.docx':
        try:
            from docx import Document
        except ImportError as e:
            raise ValueError('DOCX uploads need python-docx: pip install -r requirements.txt') from e
        try:
            doc=Document(io.BytesIO(content))
            parts=[f'[Paragraph {i}] {p.text.strip()}' for i,p in enumerate(doc.paragraphs, 1) if p.text.strip()]
            for ti,table in enumerate(doc.tables,1):
                for ri,row in enumerate(table.rows,1):
                    values=[cell.text.strip().replace('\n','; ') for cell in row.cells]
                    if any(values):parts.append(f'[Table {ti}, row {ri}] '+ ' | '.join(values))
        except Exception as e:
            raise ValueError('Unable to read DOCX') from e
    else:
        try:
            decoded=content.decode('utf-8-sig')
        except UnicodeDecodeError as e:
            raise ValueError('Text files must use UTF-8') from e
        parts=[f'[Line {i}] {line.strip()}' for i,line in enumerate(decoded.splitlines(),1) if line.strip()]
    result='\n'.join(parts)
    if not result:
        raise ValueError('No selectable text found. Scanned PDFs need OCR; paste the text instead.')
    if len(result)>49000:
        raise ValueError('Document exceeds 49,000 extracted characters; shorten it before uploading')
    return {'filename':Path(filename).name[:120], 'format':suffix[1:], 'text':result, 'source_lines':len(parts)}
